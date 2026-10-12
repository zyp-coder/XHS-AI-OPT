# -*- coding: utf-8 -*-
"""生成 功能演示.docx（功能演示.html 的 Word 版）。

用法：python3 资料/视频/生成Word.py
截图取 资料/视频/pdf素材/，缺图时留红色占位，不中断。
"""
import docx
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import os
import time

HERE = os.path.dirname(os.path.abspath(__file__))
IMG_DIR = os.path.join(HERE, "pdf素材")
OUT = os.path.join(HERE, "功能演示.docx")

RED = RGBColor(0xC4, 0x10, 0x3A)
INK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x66, 0x66, 0x66)

doc = Document()
for s in doc.sections:
    s.top_margin = Cm(1.6); s.bottom_margin = Cm(1.6)
    s.left_margin = Cm(1.8); s.right_margin = Cm(1.8)

def set_cn_font(run, size, bold=False, color=INK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"
    r = run._element.rPr.rFonts
    r.set(qn('w:eastAsia'), '微软雅黑')

def title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_cn_font(r, 28, True, RED)
    return p

def subtitle(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_cn_font(r, 16, False, INK)
    return p

def tag(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_cn_font(r, 12, False, GRAY)
    return p

def h1(text):
    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_cn_font(r, 16, True, RED)
    p.paragraph_format.space_after = Pt(6)
    return p

def body(text, size=12, bold=False, color=GRAY):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_cn_font(r, size, bold, color)
    p.paragraph_format.line_spacing = 1.5
    return p

def bullet(text, bold_head=None, size=12):
    p = doc.add_paragraph(style="List Bullet")
    if bold_head:
        r = p.add_run(bold_head)
        set_cn_font(r, size, True, INK)
        r2 = p.add_run(text)
        set_cn_font(r2, size, False, GRAY)
    else:
        r = p.add_run(text)
        set_cn_font(r, size, False, GRAY)
    p.paragraph_format.line_spacing = 1.5
    return p

def img(path, cap=None, width_cm=15.5):
    full = os.path.join(IMG_DIR, path)
    if os.path.exists(full):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(full, width=Cm(width_cm))
        if cap:
            cp = doc.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = cp.add_run(cap)
            set_cn_font(r, 9, False, RGBColor(0xA7, 0xAA, 0xB0))
    else:
        body("【截图占位：" + path + "】", 12, True, RED)

def spacer():
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

# ===== 封面 =====
title("AI雕兄")
subtitle("帮你自动「诊断 → 发笔记 → 找客户 → 私信成交」")
tag("浏览器装个插件，AI 在小红书上替你自动获客")
spacer()
for b, s in [("🔍 找客户  ", "搜笔记 · 扫评论 · 找出正有需求的潜在客户"),
             ("💬 发评论  ", "AI 生成真人话术 · 自动发 · 贴合不招黑"),
             ("📩 私信成交  ", "画像 · 定向跟进 · 把商机变成客户"),
             ("✍️ 智能生成笔记  ", "向导选题 · 成稿可改可复制 · 图文给提示词 / 视频本机出片")]:
    bullet(s, b)

# ===== 屏1 =====
h1("1. 它是干什么的 · 一条龙帮你获客")
body("AI雕兄 是一款装在浏览器里的小红书获客插件。你只需把产品资料录进「知识库」、登录自己的小红书账号，它就在你这个账号里替你干活：先诊断想清楚发什么 → 找到想要你产品的客户 → 替你在评论区发像真人一样的回复 → 用私信把围观的路人变成咨询的客户。")
img("01_笔记评论区.png", "▲ 01_笔记评论区.png · 打开任意笔记，重要评论自动高亮")
bullet("自动搜关键词相关笔记，扫评论区，挑出“主动表现出需求”的人", "🔍 找客户  ")
bullet("AI 按你的产品资料生成像真人一样的回复，自动发送，自然不招黑", "💬 发评论  ")
bullet("高意向客户进「获客清单」，AI 帮你写私信、持续跟进，直到成交", "📩 私信成交  ")
bullet("向导写笔记：卖点→人群→话题→挑稿→改稿→发稿，改稿到满意、一键复制；要图给画面表、要视频给口播稿，涨粉获客两头抓", "🧰 也能发笔记  ")

# ===== 屏2 =====
h1("2. 自动获客 · 队列「口径 × 节奏」一键开跑")
body("关键词、笔记类型、盯的账号都排进队列，点一下「开始」，它按计划一项项跑完，全程不用人工盯。")
img("02_队列_口径节奏.png", "▲ 02_队列_口径节奏.png · 队列顶栏：口径 × 节奏 + 开始 / 暂停")
bullet("看笔记就推，灌过的自动跳过，适合快速起量", "🧭 全部未灌过  ")
bullet("采样点赞变化，只挑“正在增长”的热门笔记蹭流量", "🔥 只追涨势  ")
bullet("一遍跑完，把队列一次跑到底；或定时轮，每隔 N 分钟跑一项", "⏱ 节奏自己定  ")
bullet("每轮最多处理几篇、每篇最多总评论、冷却几小时，都在设置里自己填", "🎚 跑量上限自己设  ")
img("02b_队列_运行中.png", "▲ 02b_队列_运行中.png · 正在跑 + 已完成，带计划时间 / 执行时间 / 结果")

# ===== 屏3 =====
h1("3. AI 挖商机 · 12 种评论打法")
body("打开一篇笔记，一键扫描评论区。AI 挑出真正值得回的人，直接给出可发的话术；先定这批想要什么，再按评论换招。")
img("03_挖商机_话术卡.png", "▲ 03_挖商机_话术卡.png · 商机条数 + 话术正文 + 打法标签 + 前置分析")
bullet("强制引用你的真实产品数据/个人经历，具体数字、步骤一应俱全，专业且不像广告", "📚 知识库驱动  ")
bullet("卖产品 / 品牌种草 / 主页涨粉 / 互动涨赞 / 蹭热点 / 综合自动，六选一定基调", "🎯 先定评论目的  ")
bullet("实用教程 / 避雷防坑 / 对比测评 / 数据科普 / 钩子反驳 / 评价补充 / 反向共情 / 共情共鸣 / 借势热点 / 情绪带赞 / 提问互动 / 求教自嘲——问问题的就教他，纠结的先共情，玩梗的就接梗", "🧭 12 种打法看菜下饭  ")
bullet("话术卡上标着用的是哪一招（如“获客·实用教程流”），另附一段前置分析说明为什么这么回，条条像真人一对一", "🎲 带标签带理由  ")

# ===== 屏4 =====
h1("4. 账号诊断 · 先想清楚该发什么")
body("做号最怕方向错=白努力。「账号诊断」七步走完（基础设置 / 说话风格 / 问卷 / 人设 / 产品卖点 / 知识库 / 规划），答完一次问卷就自动同步，⑦ 规划直接给你一套能照做的内容计划。")
img("08_账号诊断.png", "▲ 08_账号诊断.png · ⑦ 规划：内容规划结果 + 深度评估报告")
bullet("账号类型、讲什么内容、主要人群、建议页数与每页作用，先说什么后说什么一步排清", "📚 内容规划  ")
bullet("一周几篇、一天里几点发，跳到跳一跳够得着的节奏", "📅 发布节奏  ")
bullet("一条条能直接当标题的选题", "🎯 选题池  ")
bullet("定位 / 内容方向 / 更新节奏，先看哪个、哪里最该改", "📋 深度评估报告  ")

# ===== 屏5 =====
h1("5. 发笔记 · 从卖点一路走到出图文 / 出视频")
body("不只评论引流，还能用知识库发笔记涨粉：向导从卖点一路走到发稿——3 个候选标题、正文、5-8 个标签；改稿先停下来改到满意，再做料——做图文（出画面表 · 一句一张画面）或做视频（整段口播稿，本机出成片）。")
img("09_发笔记向导.png", "▲ 09_发笔记向导.png · 六步条（卖点 → 人群 → 话题 → 挑稿 → 改稿 → 发稿）+ ① 挑卖点")
bullet("卖点→人群→话题→挑稿→改稿→发稿一路点选，每层可换一批，发过的话题不重复", "🧭 全程点选  ")
bullet("成稿不满意就说一句「再短一点」，AI 改一版，改坏了还能改回上一版", "🤖 AI 改稿  ")
bullet("要图给画面表、要视频给口播稿——出图和发布都自己掌控", "🖼 图文 / 🎬 视频  ")

# ===== 屏6 =====
h1("6. 通知回访队列 · 别漏掉回过你的人")
body("线索 → 通知页的「通知回访队列」：把“别人回了我、我还没回”的人自动排成一条队，点「开始」一条条回访，按每日额度走，随时能停。回复内容仍走 AI 生成 + 人工确认，不打扰不硬推。")
img("04_通知回访.png", "▲ 04_通知回访.png · 通知回访队列：待回 N · 今日已回 x/y + 开始 / 停止 / 同步通知页")
bullet("谁回过你、还没回谁，自动排队，不用自己翻时间戳", "🔔 自动排队  ")
bullet("每日额度自己设，一条条走完，随时能停", "⏱ 按额度回访  ")
bullet("回复 / 定向私信一键生成，人工确认后发出", "🤖 AI 生成  ")

# ===== 屏7 =====
h1("7. 获客清单 · 私信闭环")
body("评论扫进「获客清单」→ AI 画像打标签 → 生成个性化私信 → 人工确认发送 → 定向销售。来源与画像标签一眼看清，按意向自动排序，状态分档一路跟到成交。")
img("05_获客清单.png", "▲ 05_获客清单.png · 清单表格：来源 / 画像标签 / 状态分档 / 一键联系")
bullet("来源（评论 / 点赞 / AI）与画像标签一眼看清，按意向自动排序", "🎯 画像分级  ")
bullet("待触达 / 已触达 / 谈单中 / 流失中 / 已成交 / 已废弃，一路跟到成交", "📈 状态分档跟进  ")
bullet("AI 生成个性化私信，人工确认才发", "📩 定向私信  ")

# ===== 屏8 =====
h1("8. 评论记录 · 干了多少一屏看清")
body("今天发了多少、累计多少，AI 一条条记得清清楚楚（获客 → 笔记池 → 页尾点开）。每一条评论记录——哪篇笔记、回复了谁、用了什么话术，随时回看、一键复制。")
img("06_评论记录.png", "▲ 06_评论记录.png · 折叠块摘要（今天 N 条 · 共 M 条）+ 明细清单")
bullet("今天 / 累计，清清楚楚", "💧 自动记账  ")
bullet("每一条发送记录随时回看、一键复制", "📋 可回看可复制  ")

# ===== 结尾 =====
h1("9. 别再一篇篇翻评论了")
body("让 AI 替你盯商机、替你发话术，你只负责成交。", 13, True)
body("多模型接入：DeepSeek / OpenAI / 通义 / 硅基流动，随时切换。今天就装，今天就引流。")
bullet("DeepSeek / OpenAI / 通义 / 硅基流动", "多模型接入  ")
spacer()

# ===== 联系与免责 =====
h1("10. 联系我们 · 免责声明")
body("想详细了解功能、咨询产品定制或获取授权，欢迎随时联系。", 12, False)
bullet("13360308932", "微信  ")
bullet("添加请备注「AI雕兄」", "备注  ")
doc.add_paragraph()
body("免责声明：", 13, True, RED)
bullet("本工具仅用于辅助自动化运营，请遵守小红书平台规则及相关法律法规。", "1.  ")
bullet("请务必遵守各平台社区规范，因违规操作产生的一切后果由使用者自行承担。", "2.  ")
bullet("软件功能、界面及表述不构成任何承诺，最终以实际版本为准。", "3.  ")
doc.add_paragraph()
body("© AI雕兄 · 保留所有权利", 9, False, RGBColor(0xA7, 0xAA, 0xB0))

def save_docx(document, out):
    """目标被 Word/WPS 打开着会锁住句柄（OSError errno 22），先写临时名再换名。"""
    tmp = out + ".new"
    err = None
    for _ in range(5):
        try:
            document.save(tmp)
            os.replace(tmp, out)
            return os.path.getsize(out)
        except OSError as e:
            err = e
            try:
                os.remove(tmp)
            except OSError:
                pass
            time.sleep(0.8)
    raise err


print("SAVED:", OUT, save_docx(doc, OUT), "bytes")
