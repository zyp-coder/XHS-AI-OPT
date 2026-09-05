# -*- coding: utf-8 -*-
import docx
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import os

IMG_DIR = r"D:\AIproject\小红书运营\视频\pdf素材"
OUT = r"D:\AIproject\小红书运营\视频\功能演示_优化版.docx"

RED = RGBColor(0xC4, 0x10, 0x3A)
INK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x66, 0x66, 0x66)

doc = Document()
for s in doc.sections:
    s.top_margin = Cm(1.6); s.bottom_margin = Cm(1.6)
    s.left_margin = Cm(1.8); s.right_margin = Cm(1.8)

def set_cn_font(run, size, bold=False, color=INK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"
    r = run._element.rPr.rFonts
    r.set(qn('w:eastAsia'), '微软雅黑')

def title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_cn_font(r, 28, True, RED)
    return p

def subtitle(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_cn_font(r, 16, False, INK)
    return p

def tag(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_cn_font(r, 12, False, GRAY)
    return p

def h1(text):
    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_cn_font(r, 16, True, RED)
    p.paragraph_format.space_after = Pt(6)
    return p

def body(text, size=12, bold=False, color=GRAY):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_cn_font(r, size, bold, color)
    p.paragraph_format.line_spacing = 1.5
    return p

def bullet(text, bold_head=None, size=12):
    p = doc.add_paragraph(style="List Bullet")
    if bold_head:
        r = p.add_run(bold_head)
        set_cn_font(r, size, True, INK)
        r2 = p.add_run(text)
        set_cn_font(r2, size, False, GRAY)
    else:
        r = p.add_run(text)
        set_cn_font(r, size, False, GRAY)
    p.paragraph_format.line_spacing = 1.5
    return p

def img(path, cap=None, width_cm=15.5):
    full = os.path.join(IMG_DIR, path)
    if os.path.exists(full):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(full, width=Cm(width_cm))
        if cap:
            cp = doc.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = cp.add_run(cap)
            set_cn_font(r, 9, False, RGBColor(0xA7, 0xAA, 0xB0))
    else:
        ph = body("【截图占位：" + path + "】", 12, True, RED)

def spacer():
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

# ===== 封面 =====
title("XHS智能运营工具")
subtitle("帮你自动「找客户 → 发评论 → 私信成交」")
tag("浏览器装个插件，AI 在小红书上替你自动获客")
spacer()
for b, s in [("🔍 找客户  ", "搜笔记 · 扫评论 · 找出正有需求的潜在客户"),
             ("💬 发评论  ", "AI 生成真人话术 · 自动发 · 贴合不招黑"),
             ("📩 私信成交  ", "画像 · 定向跟进 · 把商机变成客户"),
             ("✍️ 智能生成笔记  ", "AI 写文案 · 自动配图 · 一键发布")]:
    bullet(s, b)

# ===== 屏1 =====
h1("1. 它是干什么的 · 一条龙帮你获客")
body("XHS智能运营工具 是一款装在浏览器里的小红书获客插件。你只需把产品资料录进「知识库」、登录自己的小红书账号，它就在你这个账号里替你干活：找到想要你产品的客户 → 替你在评论区发像真人一样的回复 → 用私信把围观的路人变成咨询的客户。")
img("01_笔记评论区.png", "▲ 01_笔记评论区.png · 打开任意笔记，重要评论自动高亮")
bullet("自动搜关键词相关笔记，扫评论区，挑出“主动表现出需求”的人", "🔍 找客户  ")
bullet("AI 按你的产品资料生成像真人一样的回复，自动发送，自然不招黑", "💬 发评论  ")
bullet("高意向客户进「获客清单」，AI 帮你写私信、持续跟进，直到成交", "📩 私信成交  ")
bullet("AI 帮你写笔记文案、自动配图，一键发布或定时提醒，涨粉获客两头抓", "🧰 也能发笔记  ")

# ===== 屏2 =====
h1("2. 智能推广 · 两种自动策略")
body("输入关键词点开始，AI 自动搜笔记 → 找商机 → 生成话术 → 逐字发布 → 记结果，全程不用人工盯。")
img("02_冷启动_循环计划_下拉.png", "▲ 02_冷启动_循环计划_下拉.png")
bullet("看笔记就推，适合快速起量，日发 20 篇轻松跑量", "🧊 冷启动  ")
bullet("采样点赞变化，只挑“正在增长”的热门笔记蹭流量", "🔄 循环计划  ")
img("02b_智能推广_运行中.png", "▲ 02b_智能推广_运行中.png · 任务进度（成功 / 跳过 / 失败）")

# ===== 屏3 =====
h1("3. AI 挖商机 · 话术贴合真人")
body("打开一篇笔记，一键扫描评论区。AI 找出最有价值的 Top 商机，直接给出可插话的回复，三种话术随机切换。")
img("03_挖商机_话术卡.png", "▲ 03_挖商机_话术卡.png · 商机条数 + 版本 A/B/C 切换条 + 话术正文")
bullet("强制引用你的真实产品数据/个人经历，具体数字、步骤一应俱全，专业且不像广告", "📚 知识库驱动  ")
bullet("以“我最近也遇到…”亲历开场，先共情再自然带出价值，聊天语气拉近距离，最不易被认作推广", "✍️ 版本A · 朋友分享型  ")
bullet("用具体数据与行业观察把问题讲透，树立专业形象，结尾自然引导点进主页，转化更精准", "🩺 版本B · 专业顾问型  ")
bullet("像真实用户手把手分享——选哪项、点哪里、多久见效，照知识库流程一步步讲，信任感强", "📦 版本C · 产品介绍型  ")
bullet("每次随机用一种版本、三条腿轮着播，刷一圈根本看不出是批量", "🎲 三风格随机轮播  ")

# ===== 屏4 =====
h1("4. 发笔记 · AI 生成 + 定时提醒")
body("不只评论引流，还能用知识库发笔记涨粉：AI 自动生成文案 → 填好标题 / 正文 / 标签 / 配图 → 一键发布，并可设定定时提醒到点通知你该发了，持续输出不断更。")
img("07_发笔记.png", "▲ 07_发笔记.png · AI 文案生成区 + 定时提醒列表 + 一键发布")
bullet("选知识库素材 + 引导话题，自动成文，可复制或直接发布", "✍️ AI 生成文案  ")
bullet("设定提醒时间，到点通知你该发布哪一篇，保持更新节奏不用盯", "⏰ 定时提醒  ")
bullet("标题 / 正文 / 标签 / 配图全自动填好，一键确认发布", "🖼 自动配图填表  ")

# ===== 屏5 =====
h1("5. 评论跟进 · 私信或回复")
body("通知页评论跟进：把时间戳通知流按用户整理成对话。AI 判断时机，生成回复 / 定向私信，人工确认才发，不打扰不硬推。")
img("04_评论跟进.png", "▲ 04_评论跟进.png · 会话卡 + 待回复数 + AI 生成按钮")
bullet("通知流整理成对话，一眼看清谁在跟进谁", "💬 按人归组  ")
bullet("回复 / 定向私信一键生成，人工确认后发出", "🤖 AI 生成  ")

# ===== 屏6 =====
h1("6. 获客清单 · 私信闭环")
body("评论扫进「获客清单」→ AI 画像 → 生成个性化私信 → 人工确认发送 → 定向销售。高/中/低意向一目了然，成交漏斗可追踪。")
img("05_获客清单.png", "▲ 05_获客清单.png · 统计卡（总/高意向/待私信/已成交）+ 人员卡 + 意向等级")
bullet("高 / 中 / 低一目了然", "🎯 意向分级  ")
bullet("从获客到成交，每一步都有记录", "📈 漏斗可追踪  ")
bullet("AI 生成个性化私信，人工确认才发", "📩 定向私信  ")

# ===== 屏7 =====
h1("7. 运营总结 · 干了多少一屏看清")
body("今天发了多少、本周多少、累计多少，AI 一条条记得清清楚楚。每一条发布记录——哪篇笔记、回复了谁、用了什么话术，随时回看、一键复制。")
img("06_运营总结.png", "▲ 06_运营总结.png · 今天/本周/累计 三统计卡 + 按日分组发送记录")
bullet("今天 / 本周 / 累计，清清楚楚", "💧 自动记账  ")
bullet("每一条发送记录随时回看、一键复制", "📋 可回看可复制  ")

# ===== 结尾 =====
h1("8. 别再一篇篇翻评论了")
body("让 AI 替你盯商机、替你发话术，你只负责成交。", 13, True)
body("多模型接入：DeepSeek / OpenAI / 通义 / 硅基流动，随时切换。今天就装，今天就引流。")
bullet("DeepSeek / OpenAI / 通义 / 硅基流动", "多模型接入  ")
spacer()

# ===== 联系与免责 =====
h1("9. 联系我们 · 免责声明")
body("想详细了解功能、咨询产品定制或获取授权，欢迎随时联系。", 12, False)
bullet("13360308932", "微信  ")
bullet("添加请备注「XHS智能运营工具」", "备注  ")
doc.add_paragraph()
body("免责声明：", 13, True, RED)
bullet("本工具仅用于辅助自动化运营，请遵守小红书平台规则及相关法律法规。", "1.  ")
bullet("请务必遵守各平台社区规范，因违规操作产生的一切后果由使用者自行承担。", "2.  ")
bullet("软件功能、界面及表述不构成任何承诺，最终以实际版本为准。", "3.  ")
doc.add_paragraph()
body("© XHS智能运营工具 · 保留所有权利", 9, False, RGBColor(0xA7, 0xAA, 0xB0))

doc.save(OUT)
print("SAVED:", OUT, os.path.getsize(OUT), "bytes")