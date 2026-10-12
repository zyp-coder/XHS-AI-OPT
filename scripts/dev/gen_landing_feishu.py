# -*- coding: utf-8 -*-
"""生成「落地页-飞书粘贴版.docx」：从 资料/推广物料/落地页-内容稿.md 抽干净正文——
去说明块 / ✅已定 / 🔶 / 📷 标记；加微码 ×2、后台数据图、功能总览＋功能全图就位；界面截图与视频留灰字位。

跑法：py -3 scripts/dev/gen_landing_feishu.py
（成品：资料/推广物料/落地页-飞书粘贴版.docx＋备用图 资料/推广物料/飞书用图/）
"""
import re
import shutil
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "资料/推广物料/落地页-内容稿.md"
OUT = ROOT / "资料/推广物料/落地页-飞书粘贴版.docx"
IMG_DIR = ROOT / "资料/推广物料/飞书用图"
QR = ROOT / "资料/推广物料/assets/qr-wechat.png"
BACKEND = ROOT / "资料/推广物料/assets/实证-后台.jpg"
ARCH = ROOT / "资料/推广物料/assets/功能架构图.png"
OVERVIEW = ROOT / "资料/推广物料/assets/功能总览.png"

FONT = "微软雅黑"
GRAY = RGBColor(0x80, 0x80, 0x80)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
BODY = RGBColor(0x33, 0x33, 0x33)

LABELS = {"**大标题**", "**副题**", "**战绩三连**", "**按钮行**", "**小字**"}
PHOTO_RE = re.compile(r"📷〔([^〕]*)〕")
BRACE_RE = re.compile(r"〔([^〕]*)〕")
BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
NUM_RE = re.compile(r"^\d+\.\s+")

MARKER_VIDEO = "〔视频位：上传「AI雕兄-竖屏-成片-封面首帧-带BGM.mp4」（3 分钟完整演示）——传好后删掉本行〕"
MARKER_SHOT1 = "〔截图位：发笔记向导（新截图待补）——补好删掉本行〕"
MARKER_SHOT2 = "〔截图位：笔记评论区 · 商机话术卡（新截图待补）——补好删掉本行〕"
MARKER_SHOT3 = "〔截图位：获客清单 · 通知回访（新截图待补）——补好删掉本行〕"

doc = Document()


def set_font(run, size=11, bold=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def fill(p, text, size=11, color=None, bold_all=False):
    pos = 0
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            set_font(p.add_run(text[pos:m.start()]), size=size, bold=bold_all, color=color)
        set_font(p.add_run(m.group(1)), size=size, bold=True, color=color)
        pos = m.end()
    if pos < len(text):
        set_font(p.add_run(text[pos:]), size=size, bold=bold_all, color=color)
    return p


def para(style=None, align=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    return p


def gray_note(text):
    fill(para(), text, size=9, color=GRAY)


def insert_image(path, width_cm):
    p = para(align=WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    if not path.exists():
        fill(p, "〔缺图：" + path.name + "〕", size=9, color=GRAY)
        return
    p.add_run().add_picture(str(path), width=Cm(width_cm))


def handle_marker(mk):
    if ("qr-wechat" in mk) or ("加微二维码" in mk):
        insert_image(QR, 4.5)
    elif "嵌入视频" in mk:
        gray_note(MARKER_VIDEO)
    elif "发笔记向导" in mk:
        gray_note(MARKER_SHOT1)
    elif "笔记评论区" in mk:
        gray_note(MARKER_SHOT2)
    elif "获客清单" in mk:
        gray_note(MARKER_SHOT3)
    elif "实证-后台" in mk:
        insert_image(BACKEND, 9.0)
    elif "功能总览" in mk:
        insert_image(OVERVIEW, 15.2)
    elif "功能全图" in mk:
        insert_image(ARCH, 15.2)
    else:
        gray_note("〔" + mk + "〕")


def text_block(s, style=None, size=11, color=None):
    imgs = PHOTO_RE.findall(s)
    clean = PHOTO_RE.sub("", s).strip()
    clean = clean.replace("【主】", "").replace("【次】", "")
    clean = BRACE_RE.sub(lambda m: "（" + m.group(1) + "）", clean)
    if clean:
        fill(para(style=style), clean, size=size, color=color)
    for mk in imgs:
        handle_marker(mk)


def build_table(rows):
    rows_cells = []
    for r in rows:
        cs = [c.strip() for c in r.strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c or "-") for c in cs):
            continue
        rows_cells.append(cs)
    t = doc.add_table(rows=len(rows_cells), cols=2)
    t.style = "Table Grid"
    t.autofit = False
    for ri, cs in enumerate(rows_cells):
        for ci in range(2):
            cell = t.rows[ri].cells[ci]
            cell.width = Cm(3.4 if ci == 0 else 11.6)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            fill(p, cs[ci], size=10, color=BODY, bold_all=(ri == 0))
    para().paragraph_format.space_after = Pt(2)


# ---- 页面与样式 ----
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.6)
sec.top_margin = sec.bottom_margin = Cm(2.4)

normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(11)
normal.font.color.rgb = BODY
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.32
for sname, sz in (("Heading 1", 21), ("Heading 2", 15), ("Heading 3", 12.5)):
    st = doc.styles[sname]
    st.font.name = FONT
    st.font.size = Pt(sz)
    st.font.bold = True
    st.font.color.rgb = DARK
    st.paragraph_format.space_before = Pt(0 if sname == "Heading 1" else (16 if sname == "Heading 2" else 12))
    st.paragraph_format.space_after = Pt(8)
for sname in ("Normal", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Number"):
    try:
        st = doc.styles[sname]
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = rpr.makeelement(qn("w:rFonts"), {})
            rpr.insert(0, rf)
        rf.set(qn("w:eastAsia"), FONT)
    except Exception as e:  # 样式级字体只是兜底，run 级已逐个设置
        print("样式字体提示:", sname, e)

# ---- 解析底稿 ----
lines = SRC.read_text(encoding="utf-8").splitlines()
start = next(i for i, l in enumerate(lines) if l.strip() == "## 页首（第一屏）")
i, n = start + 1, len(lines)
title_pending = True

while i < n:
    s = lines[i].strip()
    i += 1
    if not s or s == "---":
        continue
    if s.startswith("> "):
        if s.startswith("> 出处"):
            gray_note(s[2:].replace("出处小字：", "出处："))
        continue  # 🔶/✅ 等内部引用块不进粘贴版
    if s.startswith("### "):
        fill(para(style="Heading 3"), s[4:].strip(), size=12.5, color=DARK, bold_all=True)
        continue
    if s.startswith("## "):
        fill(para(style="Heading 2"), s[3:].strip(), size=15, color=DARK, bold_all=True)
        continue
    if s in LABELS:
        continue
    if s.startswith("小字："):
        gray_note(s[len("小字："):])
        continue
    if s.startswith("署名："):
        gray_note(s[len("署名："):])
        continue
    if s.startswith("（数据只存本机"):
        gray_note(s)
        continue
    if s.startswith("|"):
        rows = [s]
        while i < n and lines[i].strip().startswith("|"):
            rows.append(lines[i].strip())
            i += 1
        build_table(rows)
        continue
    if s.startswith("- "):
        if s.startswith("- 【"):  # 两个动作行：不加点，正文直接写
            text_block(s[2:])
        else:
            text_block(s[2:], style="List Bullet")
        continue
    m = NUM_RE.match(s)
    if m:  # 编号保留为字面文字——List Number 样式会让多节编号连号（1..20）
        p = para()
        set_font(p.add_run(m.group(0)), size=11, bold=True, color=BODY)
        fill(p, s[m.end():], size=11, color=BODY)
        continue
    if title_pending:
        fill(para(style="Heading 1"), s, size=21, color=DARK, bold_all=True)
        title_pending = False
        continue
    text_block(s, color=BODY)

# ---- 备用图文件夹 ----
IMG_DIR.mkdir(exist_ok=True)
shutil.copy2(QR, IMG_DIR / "01-加微二维码.png")
shutil.copy2(BACKEND, IMG_DIR / "02-后台数据.jpg")
shutil.copy2(OVERVIEW, IMG_DIR / "03-功能总览.png")
shutil.copy2(ARCH, IMG_DIR / "04-功能全图.png")

doc.save(OUT)

# ---- 自检 ----
all_text = "\n".join(p.text for p in doc.paragraphs)
for t in doc.tables:
    for r in t.rows:
        for c in r.cells:
            all_text += "\n" + c.text
bad = [b for b in ("🔶", "📷", "待拍", "pdf素材", "assets/", "说明块", "您的难处", "值不值得") if b in all_text]
print("段落数:", len(doc.paragraphs), "| 图片数:", len(doc.inline_shapes))
print("灰字标记数:", all_text.count("〔"))
print("违禁残留:", bad if bad else "无")
print("输出:", OUT)
print("备用图:", IMG_DIR)
